"""Rapport complet, de l'étape 1 à l'étape 5 : sections, graphiques dessinés, texte échappé, filtre, pistes."""
import io
import unittest
from datetime import datetime, timezone
from pathlib import Path

from src.core.lecture import lire_avec_identifiants
from src.etape2_comprendre.colonnes import dictionnaire_propose, resume_colonnes
from src.etape2_comprendre.repetitions import variations
from src.etape4_visualiser.controles import graphiques_controle
from src.etape4_visualiser.interpretations import interpretations
from src.etape4_visualiser.preparation import prepare
from src.etape4_visualiser.rapport import rapport_html, rapport_markdown, synthese
from src.etape4_visualiser.rapport_word import rapport_word
from src.etape5_analyser.calculs import comparer, delais, evolution
from src.etape5_analyser.resumes import resume_comparaison, resume_delai, resume_evolution

EXEMPLE = Path(__file__).parents[1] / "fixtures" / "exemple.csv"
RESULTAT = {"total": 7, "gardees": 7, "ecartees": 0, "vides_avant": 1, "vides_apres": 1, "motifs": [],
            "moyennes": [{"colonne": "montant", "avant": 18.1, "n_avant": 5, "apres": 18.1, "n_apres": 5}]}


def contenu(**modifs):
    raw, _ = lire_avec_identifiants(EXEMPLE.read_bytes())
    dictionnaire = dictionnaire_propose(raw)
    roles = {d["Colonne"]: d["Rôle"] for d in dictionnaire}
    data, ignorees = prepare(raw, roles, False)
    fiches = resume_colonnes(raw, dictionnaire)
    base = {"nom": "exemple.csv", "taille": 1536, "ids_texte": ["id"],
            "lecture": {"separateur": "Virgule ( , )", "encodage": "UTF-8 (le plus courant)", "decimale": "Point ( 10.5 )"},
            "produit_le": datetime(2026, 10, 23, 12, 0, tzinfo=timezone.utc), "lignes": 7,
            "colonnes": 4, "empreinte": "abc123", "journal": ["Aucun changement : les données restent telles qu'importées."],
            "decisions": [{"titre": "t", "texte": "montant : 1 case vide (14,3 %)", "choix": "Laisser vide"}],
            "resultat": RESULTAT, "notes": "", "regles_verifiees": {"id : valeurs uniques": [4, 6]}, "fiches": fiches,
            "graphiques": graphiques_controle(data, roles, ignorees), "filtre": None,
            "interpretations": interpretations(fiches, {"id": variations(raw, "id")}, RESULTAT), "analyses": []}
    return {**base, **modifs}


ORDRE = ["En bref", "Étape 1 · Importer", "Étape 2 · Comprendre", "Étape 3 · Corriger", "Étape 4 · Visualiser",
         "Étape 5 · Analyser", "Ce que ce rapport ne garantit pas"]


class RapportTests(unittest.TestCase):
    def test_markdown_follows_the_five_steps(self):
        texte = rapport_markdown(contenu())
        positions = [texte.index(f"## {s}") for s in ORDRE]
        self.assertEqual(positions, sorted(positions))
        for attendu in ["- Fichier : exemple.csv (1,5 Ko).",
                        "- Lecture : séparateur Virgule ( , ), encodage UTF-8 (le plus courant), décimale Point ( 10.5 ).",
                        "- Lues comme du texte pour garder les zéros au début : id.",
                        "1 point à vérifier repéré :", "- t : 1 point",
                        "- montant : 1 case vide (14,3 %) → **Laisser vide**",
                        "| montant | Nombre | 86 % | 5 | min 10,5 · médiane 15 · moyenne 18,1 · max 30 · 1 texte |",
                        "**À retenir :**", "Graphiques de contrôle calculés sur toutes les lignes gardées.",
                        "Aucune analyse : il faut une colonne de type Catégorie ou Date"]:
            self.assertIn(attendu, texte)
        etape2, etape3 = texte.index("## Étape 2"), texte.index("## Étape 3")
        self.assertTrue(etape2 < texte.index("### Règles cochées") < etape3)
        self.assertTrue(texte.index("## Étape 4") < texte.index("### Pistes de lecture") < texte.index("## Étape 5"))

    def test_html_draws_every_chart_and_escapes_file_text(self):
        donnees = contenu(notes="<script>alert(1)</script>", nom="a<b>.csv")
        page = rapport_html(donnees)
        self.assertEqual(page.count("<svg"), len(donnees["graphiques"]) + 1)  # + le logo
        self.assertNotIn("<script>alert", page)
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", page)
        self.assertIn("Rapport qualité — a&lt;b&gt;.csv", page)
        positions = [page.index(f"<h2>{s}</h2>") for s in ORDRE]
        self.assertEqual(positions, sorted(positions))

    def test_active_filter_is_stated_in_every_version(self):
        filtre = {"texte": "categorie : exclure A", "retenues": 4, "total": 7}
        donnees = contenu(filtre=filtre)
        self.assertIn("Graphiques filtrés : categorie : exclure A (4 / 7 lignes).", synthese(donnees))
        self.assertIn("Les données téléchargées ne sont pas filtrées.", rapport_markdown(donnees))
        self.assertIn("Les données téléchargées ne sont pas filtrées.", rapport_html(donnees))


def analyses_ventes():
    """Les trois analyses de l'étape 5 sur le fichier de ventes, résumées comme à l'écran."""
    raw, _ = lire_avec_identifiants((EXEMPLE.parent / "ventes.csv").read_bytes())
    roles = {d["Colonne"]: d["Rôle"] for d in dictionnaire_propose(raw)}
    data, _ = prepare(raw, roles, True)
    data = data.assign(magasin=data["magasin"].replace({"Lille": "<b>Lille</b>"}))
    return [resume_comparaison(comparer(data, "magasin"), "magasin", "nombre", None, "lignes"),
            resume_delai(delais(data, "date_commande", "date_livraison", "magasin"), "date_commande", "date_livraison",
                         "magasin", "jours", "lignes"),
            resume_evolution(evolution(data, "date_commande"), "date_commande", None, None, "lignes")]


class AnalysesDansLeRapportTests(unittest.TestCase):
    def test_without_analyses_the_summary_says_so(self):
        self.assertIn("Aucune analyse à l'étape 5.", synthese(contenu()))

    def test_analyses_appear_in_step_5_with_table_and_findings(self):
        donnees = contenu(analyses=analyses_ventes())
        self.assertIn("3 analyses à l'étape 5.", synthese(donnees))
        texte = rapport_markdown(donnees)
        etape5 = texte.index("## Étape 5 · Analyser")
        for attendu in ["### Comparer des groupes : magasin", "### Mesurer un délai : date_commande → date_livraison",
                        "### Évolution dans le temps : date_commande", "| magasin | Nombre | Part |",
                        "**Ce que montrent les chiffres**", "Une comparaison montre une différence, pas sa cause."]:
            self.assertGreater(texte.index(attendu), etape5)

    def test_html_draws_each_analysis_and_escapes_group_names(self):
        donnees = contenu(analyses=analyses_ventes())
        page = rapport_html(donnees)
        dessinees = sum(1 for a in donnees["analyses"] if a["donnees"] is not None and len(a["donnees"]))
        self.assertEqual(page.count("<svg"), len(donnees["graphiques"]) + dessinees + 1)
        self.assertNotIn("<b>Lille</b>", page)
        self.assertIn("&lt;b&gt;Lille&lt;/b&gt;", page)


class RapportWordTests(unittest.TestCase):
    def test_word_report_follows_the_five_steps_with_every_chart(self):
        from docx import Document
        donnees = contenu(analyses=analyses_ventes(), notes="à revoir\x07")
        document = Document(io.BytesIO(rapport_word(donnees)))
        titres = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
        self.assertEqual([t for t in titres if t in ORDRE], ORDRE)
        dessinees = sum(1 for a in donnees["analyses"] if a["donnees"] is not None and len(a["donnees"]))
        self.assertEqual(len(document.inline_shapes), len(donnees["graphiques"]) + dessinees + 1)  # + le logo
        self.assertEqual(document.tables[0].rows[0].cells[0].text, "Colonne")
        self.assertIn("à revoir", [p.text for p in document.paragraphs])


class InterpretationsTests(unittest.TestCase):
    def fiche(self, **valeurs):
        return {"colonne": "x", "role": "Catégorie", "libelle": "Catégorie", "remplies": 100, "total": 100,
                "distinctes": 3, "frequences": [("A", 40), ("B", 35)], **valeurs}

    def test_balanced_data_gives_no_reading_note(self):
        self.assertEqual(interpretations([self.fiche()], {}, RESULTAT), [])

    def test_each_case_gives_one_readable_note(self):
        fiches = [self.fiche(colonne="vide", remplies=30, frequences=[("A", 10), ("B", 10)]),
                  self.fiche(colonne="zero", role="Mesure", nombres=100, moyenne=0.02, mediane=0.0),
                  self.fiche(colonne="tiree", role="Mesure", nombres=100, moyenne=30.0, mediane=10.0),
                  self.fiche(colonne="dominee", frequences=[("A", 80), ("B", 20)]),
                  self.fiche(colonne="ticket", role="Identifiant", distinctes=40, total=100)]
        repetitions = {"ticket": {"changent": [("etat", 1.0)], "identiques": []}}
        phrases = interpretations(fiches, repetitions, {**RESULTAT, "total": 100, "ecartees": 10})
        attendus = ["vide n'est remplie que pour 30 %", "zero : au moins la moitié des valeurs valent 0",
                    "tiree : la moyenne (30) est loin de la médiane (10)", "dominee : « A » représente 80 %",
                    "comptez les valeurs distinctes (40) plutôt que les lignes (100)", "10 % des lignes ont été écartées"]
        self.assertEqual(len(phrases), len(attendus))
        for phrase, attendu in zip(phrases, attendus):
            self.assertIn(attendu, phrase)


if __name__ == "__main__":
    unittest.main()
