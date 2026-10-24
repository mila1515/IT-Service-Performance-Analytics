"""Le rapport qualité en document Word (.docx), avec les graphiques en images.

Entrée : le même `contenu` que `rapport.py`. Sortie : les octets du fichier .docx.
Met en forme la liste `blocs` (mêmes sections que les autres versions, de l'étape 1 à l'étape 5) ;
les graphiques sont les images PNG des graphiques de l'écran. Les caractères de contrôle, refusés
par le format Word, sont retirés du texte issu du fichier.
"""
import io
import re

import vl_convert
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from src.etape4_visualiser.rapport import LOGO, blocs, png_graphique

SARCELLE = RGBColor(0x0B, 0x6F, 0x71)
ENCRE = RGBColor(0x1E, 0x2A, 0x38)
GRIS = RGBColor(0x5B, 0x66, 0x73)
ORANGE = RGBColor(0x9A, 0x4B, 0x00)
FOND_ENTETE = "E3F2F2"
LARGEUR_IMAGE = Cm(16)
_CONTROLE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _propre(texte):
    return _CONTROLE.sub("", str(texte))


def _document():
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Calibri", Pt(10.5), ENCRE
    for nom, taille, couleur in (("Title", 22, ENCRE), ("Heading 1", 15, SARCELLE), ("Heading 2", 12, SARCELLE)):
        style = doc.styles[nom]
        style.font.color.rgb, style.font.size, style.font.name = couleur, Pt(taille), "Calibri"
    return doc


def _paragraphe(doc, texte, couleur=None, taille=None):
    run = doc.add_paragraph().add_run(_propre(texte))
    if couleur is not None:
        run.font.color.rgb = couleur
    if taille is not None:
        run.font.size = Pt(taille)


def _tableau(doc, entetes, lignes):
    table = doc.add_table(rows=1, cols=len(entetes), style="Table Grid")
    for cellule, entete in zip(table.rows[0].cells, entetes):
        cellule.text = _propre(entete)
        cellule.paragraphs[0].runs[0].bold = True
        ombre = OxmlElement("w:shd")
        ombre.set(qn("w:val"), "clear")
        ombre.set(qn("w:fill"), FOND_ENTETE)
        cellule._tc.get_or_add_tcPr().append(ombre)
    for ligne in lignes:
        for cellule, valeur in zip(table.add_row().cells, ligne):
            cellule.text = _propre(valeur)
    for ligne in table.rows:
        for cellule in ligne.cells:
            for run in (r for p in cellule.paragraphs for r in p.runs):
                run.font.size = Pt(9)
    doc.add_paragraph()


def rapport_word(contenu):
    """Document Word : mêmes sections que le rapport HTML, graphiques compris."""
    doc = _document()
    doc.add_picture(io.BytesIO(vl_convert.svg_to_png(LOGO.read_text(encoding="utf-8"), scale=3)), width=Cm(6))
    doc.add_heading(_propre(f"Rapport qualité — {contenu['nom']}"), level=0)
    _paragraphe(doc, f"Produit le {contenu['produit_le']:%d/%m/%Y à %H:%M} (UTC).", GRIS, 9.5)
    for genre, *valeurs in blocs(contenu):
        if genre == "section":
            doc.add_heading(_propre(valeurs[0]), level=1)
        elif genre == "sous_titre":
            doc.add_heading(_propre(valeurs[0]), level=2)
        elif genre == "texte":
            _paragraphe(doc, valeurs[0])
        elif genre == "discret":
            _paragraphe(doc, valeurs[0], GRIS, 9.5)
        elif genre == "filtre":
            _paragraphe(doc, valeurs[0], ORANGE)
        elif genre == "puces":
            for element in valeurs[0]:
                doc.add_paragraph(_propre(element), style="List Bullet")
        elif genre == "decisions":
            for texte, choix in valeurs[0]:
                paragraphe = doc.add_paragraph(style="List Bullet")
                paragraphe.add_run(_propre(f"{texte} → "))
                paragraphe.add_run(_propre(choix)).bold = True
        elif genre == "tableau":
            _tableau(doc, *valeurs)
        elif genre == "graphique":
            doc.add_picture(io.BytesIO(png_graphique(valeurs[0])), width=LARGEUR_IMAGE)
        else:  # retenir
            paragraphe = doc.add_paragraph()
            paragraphe.add_run(_propre(valeurs[0])).bold = True
            paragraphe.add_run(_propre(valeurs[1]))
    tampon = io.BytesIO()
    doc.save(tampon)
    return tampon.getvalue()
