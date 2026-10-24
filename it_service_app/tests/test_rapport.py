"""Tests du rapport Word : il se génère, contient toutes ses parties et suit les options et les filtres."""
import io
import sys
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from indicateurs import calculer_kpi, filtrer  # noqa: E402
from preparation import appliquer_options, lire_csv, preparer  # noqa: E402
from rapport import rapport_word  # noqa: E402
from test_indicateurs import EXEMPLE  # noqa: E402


def texte(octets):
    doc = Document(io.BytesIO(octets))
    paragraphes = [p.text for p in doc.paragraphs]
    cellules = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join(paragraphes + cellules), doc


class Rapport(unittest.TestCase):
    def setUp(self):
        self.incidents, self.controles, self.infos = preparer(lire_csv(EXEMPLE))

    def produire(self, incidents, selection, effets=(), filtres="aucun filtre"):
        return rapport_word("incidents.csv", self.infos, self.controles, incidents, selection,
                            calculer_kpi(selection), list(effets), filtres)

    def test_toutes_les_parties(self):
        contenu, doc = texte(self.produire(self.incidents, self.incidents))
        for titre in ["1. Le fichier analysé", "Ce que représente une ligne", "Les colonnes principales",
                      "2. Contrôle et nettoyage des données", "Options de nettoyage choisies",
                      "3. Périmètre et indicateurs clés", "4. Analyses", "5. Interprétation et résumé",
                      "Limites de l'analyse"]:
            self.assertIn(titre, contenu)
        self.assertIn("66,67 %", contenu)  # KPI SLA
        self.assertIn("Aucune : les résultats suivent les règles de référence", contenu)
        self.assertGreaterEqual(len(doc.inline_shapes), 5)  # logo + graphiques

    def test_options_et_filtres_dans_le_rapport(self):
        incidents, effets = appliquer_options(self.incidents, ["sla_vide_non_conforme"])
        selection = filtrer(incidents, choix={"assignment_group": ["G1"]})
        contenu, _ = texte(self.produire(incidents, selection, effets, "Équipe = G1"))
        self.assertIn("Compter un statut SLA vide comme non conforme", contenu)
        self.assertIn("Équipe = G1", contenu)

    def test_filtre_sans_resultat(self):
        contenu, _ = texte(self.produire(self.incidents, filtrer(self.incidents, choix={"category": ["Z"]})))
        self.assertIn("Aucun incident ne correspond aux filtres", contenu)
        self.assertNotIn("4. Analyses", contenu)


if __name__ == "__main__":
    unittest.main()
