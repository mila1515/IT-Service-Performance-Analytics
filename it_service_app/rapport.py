"""Rapport d'analyse imprimable (Word .docx), de l'import du fichier jusqu'à l'interprétation.

Écrit pour un lecteur non technique : ce qu'est le fichier, ce que représente une ligne, ce que le nettoyage a fait,
les indicateurs, les graphiques de l'écran (mêmes définitions, voir graphiques.py) et un résumé.
Entrée : les objets déjà calculés par l'application. Sortie : les octets du fichier .docx.
"""
import io
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
import vl_convert
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from graphiques import (graphique_categories, graphique_delais, graphique_equipes, graphique_reassignation,
                        graphique_sla_mois, graphique_sla_priorite, graphique_volume, png)
from indicateurs import (MIN_INCIDENTS_EQUIPE, effet_reassignation, format_fr, lecture, plus_longs,
                         sla_par_priorite)
from preparation import COLONNES_UTILES

LOGO = Path(__file__).parent / "assets" / "logo.svg"
BLEU, ENCRE, GRIS_TEXTE = RGBColor(0x1E, 0x40, 0xAF), RGBColor(0x1E, 0x29, 0x3B), RGBColor(0x5B, 0x66, 0x73)
FOND_ENTETE = "E6EEFB"
LARGEUR_IMAGE = Cm(16)
_CONTROLE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

# Signification des colonnes principales d'un export d'incidents ServiceNow.
GLOSSAIRE = {
    "number": "Numéro unique de l'incident (ex. INC0000045).",
    "incident_state": "État de l'incident : New, Active, Resolved (résolu), Closed (clos)…",
    "opened_at": "Date et heure d'ouverture de l'incident.",
    "resolved_at": "Date et heure de résolution (le problème est réglé).",
    "closed_at": "Date et heure de clôture (le ticket est fermé, souvent quelques jours après la résolution).",
    "priority": "Priorité, de 1 - Critical (la plus urgente) à 4 - Low.",
    "category": "Catégorie fonctionnelle de l'incident (anonymisée dans ce jeu de données).",
    "assignment_group": "Équipe support chargée de l'incident.",
    "made_sla": "True si l'incident a respecté son engagement de délai (SLA), False sinon.",
    "reassignment_count": "Nombre de fois où l'incident a changé d'équipe.",
    "reopen_count": "Nombre de fois où l'incident a été rouvert après résolution.",
    "sys_updated_at": "Date de la mise à jour décrite par la ligne (sert à trouver le dernier état).",
    "sys_mod_count": "Nombre de modifications de l'incident au moment de la ligne.",
}

# Règles de nettoyage, expliquées simplement (toujours appliquées).
REGLES = [
    ("Une ligne par incident", "Dans un journal, chaque incident apparaît sur plusieurs lignes. On garde sa dernière "
     "version (date de mise à jour la plus récente) ; en cas d'égalité, celle qui a le plus de modifications."),
    ("Valeurs manquantes", "Les cellules vides ou contenant « ? » sont considérées comme manquantes. Une équipe, "
     "catégorie ou un lieu manquant est affiché « Inconnu » ; une personne manquante « Non assigne »."),
    ("Dates", "Les dates illisibles deviennent manquantes ; l'incident est gardé mais exclu des délais concernés."),
    ("Délais", "Délai de résolution = résolution − ouverture, en heures. Un délai négatif (résolution avant "
     "l'ouverture) est incohérent et n'est pas utilisé."),
    ("SLA", "Un statut SLA vide ou illisible n'est compté ni comme conforme ni comme non conforme : il est exclu du taux."),
    ("Aucune suppression cachée", "Seules les lignes sans numéro d'incident sont retirées, car elles ne peuvent être "
     "rattachées à aucun incident. Tout est compté dans le tableau des anomalies."),
]


def _propre(texte):
    return _CONTROLE.sub("", str(texte))


def _document():
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Calibri", Pt(10.5), ENCRE
    for nom, taille in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[nom]
        style.font.name, style.font.size = "Calibri", Pt(taille)
        style.font.color.rgb = ENCRE if nom == "Title" else BLEU
    return doc


def _paragraphe(doc, texte, couleur=None, taille=None, style=None):
    """Paragraphe ; **gras** est rendu en gras et les `code` perdent leurs accents graves."""
    paragraphe = doc.add_paragraph(style=style)
    for i, morceau in enumerate(_propre(texte).replace("`", "").split("**")):
        run = paragraphe.add_run(morceau)
        run.bold = i % 2 == 1
        if couleur is not None:
            run.font.color.rgb = couleur
        if taille is not None:
            run.font.size = Pt(taille)
    return paragraphe


def _puces(doc, lignes):
    for ligne in lignes:
        _paragraphe(doc, ligne, style="List Bullet")


def _tableau(doc, entetes, lignes, largeurs=None):
    table = doc.add_table(rows=1, cols=len(entetes), style="Table Grid")
    for cellule, entete in zip(table.rows[0].cells, entetes):
        cellule.text = _propre(entete)
        cellule.paragraphs[0].runs[0].bold = True
        ombre = OxmlElement("w:shd")
        ombre.set(qn("w:val"), "clear")
        ombre.set(qn("w:fill"), FOND_ENTETE)
        cellule._tc.get_or_add_tcPr().append(ombre)
    for ligne in lignes:
        for cellule, valeur in zip(table.add_row().cells, ligne):
            cellule.text = _propre(valeur)
    for ligne in table.rows:
        for i, cellule in enumerate(ligne.cells):
            if largeurs:
                cellule.width = Cm(largeurs[i])
            for run in (r for p in cellule.paragraphs for r in p.runs):
                run.font.size = Pt(9)
    doc.add_paragraph()


def _graphique(doc, graphique, legende):
    if graphique is None:
        return
    doc.add_picture(io.BytesIO(png(graphique)), width=LARGEUR_IMAGE)
    _paragraphe(doc, legende, GRIS_TEXTE, 9)


def _fr(valeur):
    return format_fr(valeur if not hasattr(valeur, "item") else valeur.item())


def rapport_word(nom_fichier, infos, controles, incidents, selection, kpis, effets_options, description_filtres):
    """Document Word du rapport d'analyse ; retourne ses octets."""
    doc = _document()
    doc.add_picture(io.BytesIO(vl_convert.svg_to_png(LOGO.read_text(encoding="utf-8"), scale=3)), width=Cm(2.2))
    doc.add_heading("Rapport d'analyse des incidents IT", level=0)
    _paragraphe(doc, f"Fichier analysé : {nom_fichier} · produit le {datetime.now():%d/%m/%Y à %H:%M}", GRIS_TEXTE, 9.5)

    # 1. Le fichier ------------------------------------------------------------------------------
    doc.add_heading("1. Le fichier analysé", level=1)
    _paragraphe(doc, "Ce rapport analyse un export d'incidents informatiques issu de l'outil ServiceNow. Un incident "
                "est une demande d'aide ou une panne signalée au support IT (poste qui ne démarre plus, application "
                "inaccessible…). L'objectif est de mesurer la qualité du service : volumes traités, délais de "
                "résolution, respect des engagements de délai (SLA) et réassignations entre équipes.")
    debut, fin = infos["periode"]
    if infos["format"] == "journal d'événements":
        _paragraphe(doc, f"**Ce que représente une ligne.** Le fichier est un journal : chaque ligne décrit une "
                    f"mise à jour d'un incident (ouverture, changement d'équipe, résolution, clôture…). Un même "
                    f"incident apparaît donc sur plusieurs lignes, {_fr(infos['lignes_par_incident'])} en moyenne. "
                    f"Pour compter les incidents, on ne garde que la dernière version de chacun.")
    else:
        _paragraphe(doc, "**Ce que représente une ligne.** Le fichier est une table d'incidents : chaque ligne "
                    "correspond à un incident, dans son dernier état connu.")
    _tableau(doc, ["Élément", "Valeur"], [
        ["Lignes du fichier", _fr(infos["lignes"])], ["Colonnes", _fr(infos["colonnes"])],
        ["Incidents différents", _fr(infos["incidents"])],
        ["Période couverte (ouvertures)", "—" if pd.isna(debut) else f"{debut:%d/%m/%Y} au {fin:%d/%m/%Y}"]], largeurs=[7, 9])
    doc.add_heading("Les colonnes principales", level=2)
    presentes = [c for c in GLOSSAIRE if c not in infos["colonnes_absentes"]]
    _tableau(doc, ["Colonne", "Signification"], [[c, GLOSSAIRE[c]] for c in presentes], largeurs=[4, 12])
    if infos["colonnes_absentes"]:
        _paragraphe(doc, "Colonnes absentes du fichier (les indicateurs qui en dépendent sont indisponibles) : "
                    + ", ".join(f"{c} ({COLONNES_UTILES[c]})" for c in infos["colonnes_absentes"]) + ".")

    # 2. Nettoyage -------------------------------------------------------------------------------
    doc.add_heading("2. Contrôle et nettoyage des données", level=1)
    _paragraphe(doc, "Avant tout calcul, le fichier est contrôlé et préparé selon des règles fixes, identiques "
                "d'un fichier à l'autre :")
    for titre, texte in REGLES:
        _paragraphe(doc, f"**{titre}.** {texte}", style="List Bullet")
    anomalies = controles[controles["valeur"] > 0]
    doc.add_heading("Ce qui a été trouvé dans ce fichier", level=2)
    if anomalies.empty:
        _paragraphe(doc, "Aucune anomalie détectée.")
    else:
        _tableau(doc, ["Contrôle", "Nombre", "Effet sur les calculs"],
                 [[r["contrôle"], _fr(r["valeur"]), r["effet sur les calculs"]] for _, r in anomalies.iterrows()],
                 largeurs=[7, 2.5, 6.5])
    doc.add_heading("Options de nettoyage choisies", level=2)
    if effets_options:
        _tableau(doc, ["Option", "Incidents concernés", "Effet"],
                 [[e["option"], _fr(e["incidents concernés"]), e["effet"]] for e in effets_options],
                 largeurs=[6, 2.5, 7.5])
    else:
        _paragraphe(doc, "Aucune : les résultats suivent les règles de référence, identiques au rapport Power BI.")

    # 3. Périmètre et indicateurs ----------------------------------------------------------------
    doc.add_heading("3. Périmètre et indicateurs clés", level=1)
    _paragraphe(doc, f"**Incidents analysés :** {_fr(selection['number'].nunique())} sur {_fr(len(incidents))} · "
                f"**Filtres :** {description_filtres}.")
    _tableau(doc, ["Indicateur", "Valeur", "Ce qu'il mesure", "Calculé sur"],
             [[k["indicateur"], format_fr(k["valeur"], k["unité"]) if not k["indisponible"] else "indisponible",
               k["définition"] if not k["indisponible"] else k["indisponible"], k["population"]] for k in kpis],
             largeurs=[3.6, 2.4, 5.5, 4.5])

    if selection.empty:
        _paragraphe(doc, "Aucun incident ne correspond aux filtres : pas d'analyse détaillée.")
        return _octets(doc)

    # 4. Analyses --------------------------------------------------------------------------------
    doc.add_heading("4. Analyses", level=1)
    doc.add_heading("Évolution dans le temps", level=2)
    _graphique(doc, graphique_volume(selection), "Nombre d'incidents par mois d'ouverture. Un mois avec peu "
               "d'incidents ne permet pas de conclure à une tendance.")
    if selection["sla_compliant"].notna().any():
        _graphique(doc, graphique_sla_mois(selection), "Taux de respect des SLA par mois d'ouverture.")

    doc.add_heading("Respect des SLA par priorité", level=2)
    priorites = sla_par_priorite(selection)
    if not priorites.empty:
        _graphique(doc, graphique_sla_priorite(selection), "Part des incidents qui respectent leur SLA, par priorité.")
        _tableau(doc, ["Priorité", "Incidents", "SLA (%)", "Délai moyen (h)", "Délai médian (h)"],
                 [[p, _fr(n), _fr(s), _fr(m), _fr(md)] for p, n, s, m, md in priorites.itertuples(index=False)])

    doc.add_heading("Effet des réassignations", level=2)
    effet = effet_reassignation(selection)
    if not effet.empty:
        _graphique(doc, graphique_reassignation(selection), "Délai moyen de résolution selon que l'incident a "
                   "changé d'équipe ou non. La comparaison montre un lien, pas une cause.")
        _tableau(doc, ["Réassignation", "Incidents", "SLA (%)", "Délai moyen (h)", "Délai médian (h)"],
                 [[r, _fr(n), _fr(s), _fr(m), _fr(md)] for r, n, s, m, md in effet.itertuples(index=False)])

    doc.add_heading("Délais de résolution", level=2)
    _graphique(doc, graphique_delais(selection), "Nombre d'incidents par tranche de délai. La moyenne est tirée vers "
               "le haut par quelques incidents très longs : la médiane est plus représentative.")

    doc.add_heading("Équipes et catégories", level=2)
    _graphique(doc, graphique_equipes(selection), f"Respect des SLA des équipes les plus chargées (au moins "
               f"{MIN_INCIDENTS_EQUIPE} incidents). En gris : incidents sans équipe renseignée.")
    _graphique(doc, graphique_categories(selection), "Les 15 catégories les plus fréquentes ; en gris, les "
               "catégories non renseignées et les suivantes regroupées.")

    longs = plus_longs(selection, top=10)
    if not longs.empty:
        doc.add_heading("Incidents les plus longs", level=2)
        _tableau(doc, ["Incident", "Délai (h)", "Équipe", "Catégorie", "Priorité"],
                 [[r.get("number"), _fr(r.get("resolution_time_hours")), r.get("assignment_group", ""),
                   r.get("category", ""), r.get("priority", "")] for _, r in longs.iterrows()])

    # 5. Interprétation --------------------------------------------------------------------------
    doc.add_heading("5. Interprétation et résumé", level=1)
    messages = lecture(selection, kpis)
    _puces(doc, messages or ["Pas assez de données pour formuler des constats."])
    doc.add_heading("Limites de l'analyse", level=2)
    _puces(doc, ["Les comparaisons montrent des liens entre variables, pas des causes : elles orientent les "
                 "vérifications sans les remplacer.",
                 "Les délais sont calendaires (24 h sur 24), sans tenir compte des heures ouvrées ni des jours fériés.",
                 "Le respect des SLA repose sur le statut fourni par ServiceNow ; les objectifs de délai eux-mêmes "
                 "ne figurent pas dans le fichier."])
    return _octets(doc)


def _octets(doc):
    tampon = io.BytesIO()
    doc.save(tampon)
    return tampon.getvalue()
